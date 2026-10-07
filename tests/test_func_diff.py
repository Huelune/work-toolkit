import argparse
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
import func_diff as fdf  # noqa: E402


def names(src):
    return fdf.find_functions(src.splitlines())


def make_tree(root: Path, files: dict):
    for rel, data in files.items():
        p = root / rel
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(data, encoding="utf-8")


def ns(**kw):
    base = dict(exclude=None, ext=[".c"], ignore_whitespace=False, ignore_blank_lines=False, context=3)
    base.update(kw)
    return argparse.Namespace(**base)


A_MOTOR = ("#define GAIN 2\n"
           "static int g;\n"
           "void Motor_Init(void)\n{\n  g = 0;\n}\n"
           "void Motor_Step(void)\n{\n  g += GAIN;\n}\n"
           "void Motor_Old(void)\n{\n}\n")
B_MOTOR = ("#define GAIN 3\n"
           "static int g;\n"
           "void Motor_Step(void)\n{\n  g += GAIN;\n}\n"
           "void Motor_Init(void)\n{\n  g = 1;\n}\n"
           "void Motor_New(void)\n{\n}\n")
DUP = "void f(void) {\n}\n#if X\nvoid f(void) {\n}\n#endif\n"


class TreeMixin:
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.a, self.b = Path(self.tmp.name, "A"), Path(self.tmp.name, "B")
        make_tree(self.a, {"src/motor.c": A_MOTOR, "bad.c": "void f(void) {\n", "dup.c": DUP,
                           "a_only.c": "void x(void)\n{\n}\n", "readme.txt": "x"})
        make_tree(self.b, {"src/motor.c": B_MOTOR, "bad.c": "void f(void) {\n}\n", "dup.c": DUP})

    def tearDown(self):
        self.tmp.cleanup()


class ScannerTest(unittest.TestCase):
    def test_plain_and_static(self):
        src = "int g;\n\nint add(int a, int b)\n{\n  return a + b;\n}\nstatic void helper(void) { }\n"
        self.assertEqual(names(src), [("add", 3, 6), ("helper", 7, 7)])

    def test_simulink_style(self):
        src = ("/* Model step function */\n"
               "void Model_step(void)\n"
               "{\n"
               "  if (x) {\n"
               "    y = 1;\n"
               "  }\n"
               "}\n")
        self.assertEqual(names(src), [("Model_step", 2, 7)])

    def test_macro_and_extension_signatures(self):
        src = ("FUNC(void, RTE_CODE) Rte_Foo(void)\n{\n}\n"
               "__interrupt void ISR_Timer(void) {\n}\n")
        self.assertEqual(names(src), [("Rte_Foo", 1, 3), ("ISR_Timer", 4, 5)])

    def test_ifdef_split_signature(self):
        src = "#ifdef A\nvoid f(int a)\n#else\nvoid f(void)\n#endif\n{\n}\n"
        self.assertEqual(names(src), [("f", 2, 7)])

    def test_truncated_literal_at_eof(self):
        self.assertEqual(fdf.mask_c("x = '\\"), "x =   ")
        self.assertEqual(fdf.mask_c('x = "\\'), "x =   ")
        self.assertEqual(names("void f(void)\n{\n}\nx = '\\"), [("f", 1, 3)])

    def test_not_functions(self):
        src = ("void proto(void);\n"
               "void (*fp)(int);\n"
               "int arr[] = { 1, 2 };\n"
               "struct S { int a; } s = { 0 };\n"
               "typedef struct { int b; } T;\n"
               "struct P __attribute__((packed)) { int c; };\n"
               "enum E { A, B };\n")
        self.assertEqual(names(src), [])

    def test_braces_in_comments_strings_and_macros(self):
        src = ('#define BLOCK() do { \\\n'
               '    x(); } while (0)\n'
               'void f(int a /* { */)\n'
               '{\n'
               '  const char *s = "}\\"{";\n'
               "  char c = '}';\n"
               '  // }\n'
               '}\n')
        self.assertEqual(names(src), [("f", 3, 8)])

    def test_unbalanced_returns_none(self):
        self.assertIsNone(names("#ifdef A\nvoid f(int a) {\n#else\nvoid f(void) {\n#endif\n}\n"))
        self.assertIsNone(names("void f(void) {\n}\n}\n"))

    def test_split_units(self):
        src = "int g;\nvoid f(void) {\n}\n#if A\nvoid f(void) {\n}\n#endif\n"
        u = fdf.split_units(src.splitlines())
        self.assertEqual(list(u), [fdf.FUNC_OUTSIDE, "f", "f#2"])
        self.assertEqual(u[fdf.FUNC_OUTSIDE], [1, 4, 7])
        self.assertEqual(u["f"], [2, 3])
        self.assertEqual(u["f#2"], [5, 6])

    def test_comment_continues_from_directive(self):
        src = "#define X 1 /* a\n b */\nvoid f(void)\n{\n}\n"
        self.assertEqual(names(src), [("f", 3, 5)])

    def test_trailing_attribute(self):
        src = 'void f(void) __attribute__((section(".x")))\n{\n}\nint g(int a) __declspec(noinline)\n{\n}\n'
        self.assertEqual(names(src), [("f", 1, 3), ("g", 4, 6)])

    def test_osek_task_isr_macros(self):
        self.assertEqual(names("TASK(T10ms)\n{\n}\nISR(Can_Isr) {\n}\nTASK(T20ms)\n{\n}\n"),
                         [("T10ms", 1, 3), ("Can_Isr", 4, 5), ("T20ms", 6, 8)])
        self.assertEqual(names("void ISR(void)\n{\n}\nmain(void)\n{\n}\n"), [("ISR", 1, 3), ("main", 4, 6)])
        self.assertEqual(names("FUNC(void, CODE) Rte_Foo(void)\n{\n}\n"), [("Rte_Foo", 1, 3)])

    def test_directive_comment_then_continuation(self):
        self.assertEqual(names("#define M(x) /* c */ \\\n  do { x } while(0)\nvoid f(void)\n{\n}\n"), [("f", 3, 5)])
        self.assertEqual(names("#define M(x) /* c */ {\nvoid f(void)\n{\n}\n"), [("f", 2, 4)])

    def test_comment_markers_inside_directive_strings(self):
        src = ('#define S "/*"\nvoid f(void)\n{\n}\n'
               '#endif // see /* x\nvoid g(void)\n{\n}\n'
               '#include "a/*b"\nvoid h(void)\n{\n}\n')
        self.assertEqual(names(src), [("f", 2, 4), ("g", 6, 8), ("h", 10, 12)])

    def test_unterminated_quotes(self):
        self.assertEqual(names("#error don't do this\nvoid f(void)\n{\n}\n"), [("f", 2, 4)])
        self.assertEqual(names("#if 0\nx = 'a\n#define X {\n#endif\nint y;\nvoid g(void)\n{\n}\n"), [("g", 6, 8)])


class ListTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()

    def tearDown(self):
        self.tmp.cleanup()

    def test_excel_layout(self):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        for row in [("파일명", "함수명"), ("motor.c", "Motor_Init"), (None, "Motor_Step"),
                    ("sensor.c", "Sensor_Read"), ("sensor.c", "(함수 외)"), (None, None),
                    ("util.c", None), ("motor.c", "Motor_Stop")]:
            ws.append(row)
        ws.merge_cells("A2:A3")
        p = Path(self.tmp.name, "list.xlsx")
        wb.save(p)
        self.assertEqual(fdf.load_func_list(p), {
            "motor.c": {"Motor_Init", "Motor_Step", "Motor_Stop"},
            "sensor.c": {"Sensor_Read", fdf.FUNC_OUTSIDE},
            "util.c": None})

    def test_excel_header_after_blank_first_row(self):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append([None, None])
        ws.append(["파일명", "함수명"])
        ws.append(["motor.c", "Motor_Init"])
        p = Path(self.tmp.name, "list.xlsx")
        wb.save(p)
        self.assertEqual(fdf.load_func_list(p), {"motor.c": {"Motor_Init"}})

    def test_text_layout(self):
        p = Path(self.tmp.name, "list.txt")
        p.write_text("파일,함수\nmotor.c,Motor_Init\n,Motor_Step\n# 주석\nutil.c\nutil.c,Ignored\n",
                     encoding="utf-8")
        self.assertEqual(fdf.load_func_list(p), {"motor.c": {"Motor_Init", "Motor_Step"}, "util.c": None})


class CompareTest(TreeMixin, unittest.TestCase):
    def result(self, plan=None, **kw):
        res = fdf.compare_funcs(self.a, self.b, ns(**kw), plan)
        return [(r["file"], r["func"], r["status"]) for r in res["rows"]]

    def test_all_functions(self):
        self.assertEqual(self.result(), [
            ("a_only.c", "x", fdf.ONLY_A),
            ("bad.c", fdf.WHOLE_FILE, fdf.BROKEN),
            ("dup.c", fdf.FUNC_OUTSIDE, fdf.SAME),
            ("dup.c", "f", fdf.SAME),
            ("dup.c", "f#2", fdf.SAME),
            ("src/motor.c", fdf.FUNC_OUTSIDE, fdf.CHANGED),
            ("src/motor.c", "Motor_Init", fdf.CHANGED),
            ("src/motor.c", "Motor_Step", fdf.SAME),
            ("src/motor.c", "Motor_Old", fdf.ONLY_A),
            ("src/motor.c", "Motor_New", fdf.ONLY_B)])

    def test_row_details(self):
        rows = fdf.compare_funcs(self.a, self.b, ns())["rows"]
        init = next(r for r in rows if r["func"] == "Motor_Init")
        self.assertEqual((init["a_range"], init["b_range"], init["added"], init["removed"]), ("3-6", "7-10", 1, 1))
        old = next(r for r in rows if r["func"] == "Motor_Old")
        self.assertEqual((old["a_range"], old["b_range"], old["removed"]), ("11-13", "", 3))

    def test_plan(self):
        plan = {"motor.c": {"Motor_Step", "motor_init", "Motor_Old"}, "Motor.C": {"Motor_New"},
                "nofile.c": {"X"}, "dup.c": {"f"}}
        self.assertEqual(self.result(plan), [
            ("nofile.c", "X", fdf.NO_FILE),
            ("dup.c", "f", fdf.SAME),
            ("dup.c", "f#2", fdf.SAME),
            ("src/motor.c", "Motor_Step", fdf.SAME),
            ("src/motor.c", "Motor_Old", fdf.ONLY_A),
            ("src/motor.c", "Motor_New", fdf.ONLY_B),
            ("src/motor.c", "motor_init", fdf.NO_FUNC)])

    def test_numbered_entry_in_list(self):
        self.assertEqual(self.result({"dup.c": {"f#2"}}), [("dup.c", "f#2", fdf.SAME)])
        self.assertEqual(self.result({"dup.c": {"f", "f#2"}}),
                         [("dup.c", "f", fdf.SAME), ("dup.c", "f#2", fdf.SAME)])

    def test_blank_function_cell_excludes_outside(self):
        funcs = [("src/motor.c", f, s) for f, s in (("Motor_Init", fdf.CHANGED), ("Motor_Step", fdf.SAME),
                                                     ("Motor_Old", fdf.ONLY_A), ("Motor_New", fdf.ONLY_B))]
        self.assertEqual(self.result({"src/motor.c": None}), funcs)
        outside = [("src/motor.c", fdf.FUNC_OUTSIDE, fdf.CHANGED)]
        self.assertEqual(self.result({"motor.c": None, "Motor.c": {fdf.FUNC_OUTSIDE}}), outside + funcs)
        self.assertEqual(self.result({"Motor.c": {fdf.FUNC_OUTSIDE}, "motor.c": None}), outside + funcs)

    def test_many_functions_is_fast(self):
        import time
        n = 6000
        body = "".join("void f%d(void)\n{\n  x = %d;\n}\n" % (i, i) for i in range(n))
        make_tree(self.a, {"big/big.c": body})
        make_tree(self.b, {"big/big.c": body.replace("x = 7;", "x = 8;").replace("x = 1234;", "x = 0;")})
        t = time.perf_counter()
        res = fdf.compare_funcs(self.a / "big", self.b / "big", ns())
        elapsed = time.perf_counter() - t
        print(f"[perf] {n} functions: {elapsed:.2f}s")
        self.assertEqual(len(res["rows"]), n)
        self.assertEqual(sum(r["status"] == fdf.CHANGED for r in res["rows"]), 2)
        self.assertLess(elapsed, 3.0)

    def test_ignore_whitespace(self):
        make_tree(self.a, {"w.c": "void f(void)\n{\n  x();\n}\n"})
        make_tree(self.b, {"w.c": "void f(void)\n{\n\tx();   \n}\n"})
        plan = {"w.c": {"f"}}
        self.assertEqual(self.result(plan), [("w.c", "f", fdf.CHANGED)])
        self.assertEqual(self.result(plan, ignore_whitespace=True), [("w.c", "f", fdf.SAME)])


def run_cli(*argv):
    return subprocess.run([sys.executable, str(ROOT / "tools" / "func_diff.py"), *map(str, argv), "--no-color"],
                          capture_output=True, text=True, encoding="utf-8", errors="replace",
                          env={**os.environ, "PYTHONIOENCODING": "utf-8"})  # 파이프 stderr 인코딩을 고정


class CliTest(TreeMixin, unittest.TestCase):
    def test_reports_and_exit_code(self):
        out = Path(self.tmp.name, "out", "rep")
        r = run_cli(self.a, self.b, "-o", out)
        self.assertEqual(r.returncode, 2, r.stderr)  # bad.c 분석 불가
        self.assertIn("Motor_Init", r.stdout)
        self.assertIn("g = 1;", r.stdout)  # 변경 함수 diff 출력
        import openpyxl
        ws = openpyxl.load_workbook(out.with_suffix(".xlsx"))["함수 목록"]
        self.assertEqual([c.value for c in ws[1]][-2:], ["검토 결과", "검토 의견"])
        values = [tuple(c.value for c in row[:3]) for row in ws.iter_rows(min_row=2)]
        self.assertIn(("src/motor.c", "Motor_Init", "변경"), values)
        cell = next(row[1] for row in ws.iter_rows(min_row=2) if row[1].value == "Motor_Init")
        self.assertEqual(cell.hyperlink.target, "rep.html")
        page = out.with_suffix(".html").read_text(encoding="utf-8")
        self.assertIn(f"id='{cell.hyperlink.location}'", page)

    def test_list_mode_exit_codes(self):
        lst = Path(self.tmp.name, "list.txt")
        lst.write_text("motor.c,Motor_Step\n", encoding="utf-8")
        self.assertEqual(run_cli(self.a, self.b, "-l", lst, "--summary").returncode, 0)
        lst.write_text("motor.c,Motor_Init\n", encoding="utf-8")
        self.assertEqual(run_cli(self.a, self.b, "-l", lst, "--summary").returncode, 1)
        lst.write_text("파일명,함수명\n", encoding="utf-8")
        r = run_cli(self.a, self.b, "-l", lst)
        self.assertNotEqual(r.returncode, 0)
        self.assertIn("비어 있습니다", r.stderr)


if __name__ == "__main__":
    unittest.main()
