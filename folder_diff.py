#!/usr/bin/env python3
"""
두 폴더 비교 도구 (folder_diff.py)

같은 상대 경로(하위 폴더 포함)의 파일끼리 짝지어 비교하고,
변경된 파일은 어느 줄이 어떻게 바뀌었는지 보여줍니다.

사용 예:
    python folder_diff.py 폴더A 폴더B
    python folder_diff.py 폴더A 폴더B -o report      # report.xlsx + report.html
    python folder_diff.py 폴더A 폴더B -w --exclude build "src/*.tmp"
    python folder_diff.py 폴더A 폴더B --ext .py .c .h .m
    python folder_diff.py 폴더A 폴더B -l 대상목록.xlsx -o report

옵션:
    -w, --ignore-whitespace   줄 앞뒤 공백/줄바꿈(CRLF·LF) 차이 무시
    -B, --ignore-blank-lines  빈 줄 차이 무시
    -x, --exclude PATTERN...  제외할 파일/폴더 패턴 (glob, '/'가 들어가면 상대 경로 전체와 비교)
    -e, --ext EXT...          지정한 확장자만 비교
    -l, --list FILE           목록에 있는 파일만 비교 (.xlsx는 첫 시트 A열, .txt는 한 줄에 하나)
                              파일명만 쓰면 하위 폴더 어디에 있든 찾음 (대소문자 무시, 경로·glob도 가능)
    -c, --context N           diff 앞뒤로 보여줄 줄 수 (기본 3)
    -o, --report NAME         NAME.xlsx(파일 목록)와 NAME.html(변경 내용)을 함께 저장
    --html FILE               변경 내용을 나란히 비교하는 HTML 리포트만 저장
    --excel FILE              파일 목록 엑셀(.xlsx) 리포트만 저장 (openpyxl 필요: pip install openpyxl)
    --summary                 파일 목록 요약만 출력 (diff 내용 생략)
    --no-color                콘솔 색상 끄기

종료 코드: 0 = 차이 없음, 1 = 차이 있음, 2 = 읽기 오류 발생
엑셀 리포트 외에는 표준 라이브러리만 사용하므로 별도 설치가 필요 없습니다 (Python 3.8+).
"""
import argparse
import codecs
import difflib
import fnmatch
import hashlib
import html
import os
import sys
from datetime import datetime
from pathlib import Path

DEFAULT_EXCLUDES = [".git", ".svn", "__pycache__", ".idea", ".vscode", "*.pyc", "*.asv", "slprj"]
# cp949는 euc-kr을 포함하고, latin-1은 어떤 바이트든 디코딩되므로 마지막 대체용
ENCODINGS = ["utf-8-sig", "cp949", "latin-1"]
UTF16_BOMS = (codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)
INLINE_DIFF_MAX = 2000  # 이보다 긴 줄은 HTML에서 글자 단위 강조 생략
# 비교 대상 목록의 첫 칸이 이 중 하나면 제목으로 보고 건너뜀
TARGET_HEADERS = {"파일", "파일명", "파일 이름", "파일이름", "경로", "file", "filename", "file name", "name", "path"}


# ---------------------------------------------------------------- 파일 수집
def is_excluded(rel: Path, patterns) -> bool:
    path = rel.as_posix()
    for p in patterns:
        if "/" in p:
            if fnmatch.fnmatch(path, p):
                return True
        elif any(fnmatch.fnmatch(part, p) for part in rel.parts):
            return True
    return False


def collect_files(root: Path, excludes, exts, errors):
    files = {}

    def on_error(ex):
        errors.append({"path": str(ex.filename), "detail": f"폴더 읽기 실패: {ex.strerror or ex}"})

    for dirpath, dirnames, filenames in os.walk(root, onerror=on_error):
        rel_dir = Path(dirpath).relative_to(root)
        # 제외 폴더는 하위로 내려가지 않음
        dirnames[:] = [d for d in dirnames if not is_excluded(rel_dir / d, excludes)]
        for name in filenames:
            rel = rel_dir / name
            if is_excluded(rel, excludes):
                continue
            if exts and Path(name).suffix.lower() not in exts:
                continue
            files[rel.as_posix()] = Path(dirpath) / name
    return files


# ---------------------------------------------------------------- 파일 읽기
def file_hash(path: Path) -> str:
    h = hashlib.md5()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def is_binary(data: bytes) -> bool:
    return not data.startswith(UTF16_BOMS) and b"\x00" in data[:8192]


def decode_lines(data: bytes):
    """줄 목록과 파일 형식 정보(줄바꿈·인코딩·BOM·끝 줄바꿈)를 반환"""
    if data.startswith(UTF16_BOMS):
        text, enc = data.decode("utf-16"), "utf-16"
    else:
        for enc in ENCODINGS:
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
    if text.startswith("﻿"):  # utf-16 디코딩 시 남는 BOM 제거
        text = text[1:]
    crlf = text.count("\r\n")
    lf = text.count("\n") - crlf
    eol = "CRLF" if crlf and not lf else "LF" if lf and not crlf else "혼합" if crlf else "없음"
    fmt = {"줄바꿈": eol,
           "인코딩": enc.replace("utf-8-sig", "utf-8"),  # BOM 여부는 따로 표시
           "BOM": "있음" if data.startswith((codecs.BOM_UTF8,) + UTF16_BOMS) else "없음",
           "끝 줄바꿈": "있음" if text.endswith("\n") else "없음"}
    return text.splitlines(), fmt


def normalize(lines, ignore_ws, ignore_blank):
    """(원래 줄 번호, 원문, 비교용 키) 목록. 화면에는 원문과 원래 줄 번호를 보여줌"""
    out = []
    for no, line in enumerate(lines, 1):
        if ignore_blank and not line.strip():
            continue
        out.append((no, line, line.strip() if ignore_ws else line))
    return out


# ---------------------------------------------------------------- 비교 대상 목록
def load_targets(path: Path):
    """엑셀(.xlsx) 첫 시트의 A열, 또는 텍스트 파일의 각 줄에서 비교 대상 목록을 읽음"""
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        values = [row[0] for row in wb.worksheets[0].iter_rows(max_col=1, values_only=True) if row]
        wb.close()
    else:
        values = [l for l in decode_lines(path.read_bytes())[0] if not l.lstrip().startswith("#")]
    values = [str(v).strip() for v in values if v is not None and str(v).strip()]
    if values and values[0].casefold() in TARGET_HEADERS:  # 첫 칸이 제목이면 건너뜀
        values = values[1:]
    return list(dict.fromkeys(values))  # 중복 제거 (순서 유지)


def match_targets(files, targets):
    """목록 항목(파일명, 상대 경로 또는 glob 패턴, 대소문자 무시)에 맞는 파일과,
    양쪽 폴더 어디에서도 찾지 못한 항목을 반환"""
    by_name, by_path = {}, {}
    for rel in files:
        by_name.setdefault(rel.rsplit("/", 1)[-1].casefold(), []).append(rel)
        by_path[rel.casefold()] = rel
    selected, missing = set(), []
    for t in targets:
        key = t.replace("\\", "/").strip("/").casefold()
        if any(ch in key for ch in "*?["):
            hits = [r for r in files if fnmatch.fnmatchcase(
                r.casefold() if "/" in key else r.rsplit("/", 1)[-1].casefold(), key)]
        elif "/" in key:
            hits = [by_path[key]] if key in by_path else []
        else:
            hits = by_name.get(key, [])
        selected.update(hits)
        if not hits:
            missing.append(t)
    return selected, missing


# ---------------------------------------------------------------- 비교
def compare(dir_a: Path, dir_b: Path, args, targets=None):
    excludes = [p.replace("\\", "/").strip("/") for p in DEFAULT_EXCLUDES + (args.exclude or [])]
    exts = {e.lower() if e.startswith(".") else "." + e.lower() for e in (args.ext or [])}
    errors = []
    fa = collect_files(dir_a, excludes, exts, errors)
    fb = collect_files(dir_b, excludes, exts, errors)
    missing = []
    if targets is not None:  # 목록에 있는 파일만 비교
        selected, missing = match_targets(set(fa) | set(fb), targets)
        fa = {r: p for r, p in fa.items() if r in selected}
        fb = {r: p for r, p in fb.items() if r in selected}

    result = {"only_a": sorted(set(fa) - set(fb)),
              "only_b": sorted(set(fb) - set(fa)),
              "same": [], "changed": [], "binary_changed": [],
              "format_changed": [],  # 내용은 같고 줄바꿈/인코딩/BOM 등 형식만 다름
              "ignored": [],         # -w/-B 옵션으로 무시된 공백 차이만 있음
              "missing": missing,    # 비교 대상 목록에는 있지만 양쪽 폴더 모두에 없음
              "errors": errors}      # 권한 없음·잠긴 파일 등 읽기 실패

    for rel in sorted(set(fa) & set(fb)):
        try:
            compare_file(rel, fa[rel], fb[rel], args, result)
        except OSError as ex:
            errors.append({"path": rel, "detail": f"파일 읽기 실패: {ex.strerror or ex}"})
    return result


def compare_file(rel, pa: Path, pb: Path, args, result):
    if pa.stat().st_size == pb.stat().st_size and file_hash(pa) == file_hash(pb):
        result["same"].append(rel)
        return
    da, db = pa.read_bytes(), pb.read_bytes()
    if is_binary(da) or is_binary(db):
        result["binary_changed"].append(rel)
        return
    la, fmt_a = decode_lines(da)
    lb, fmt_b = decode_lines(db)
    na = normalize(la, args.ignore_whitespace, args.ignore_blank_lines)
    nb = normalize(lb, args.ignore_whitespace, args.ignore_blank_lines)
    ka, kb = [x[2] for x in na], [x[2] for x in nb]
    if ka == kb:  # 줄 내용은 같은데 바이트가 다른 경우
        ws_keys = ("줄바꿈", "끝 줄바꿈") if args.ignore_whitespace else ()
        detail = [f"{k} {fmt_a[k]}→{fmt_b[k]}" for k in fmt_a
                  if fmt_a[k] != fmt_b[k] and k not in ws_keys]
        if detail:
            result["format_changed"].append({"path": rel, "detail": ", ".join(detail)})
        else:
            result["ignored"].append(rel)
        return
    sm = difflib.SequenceMatcher(None, ka, kb)
    ops = sm.get_opcodes()
    result["changed"].append({
        "path": rel, "a": na, "b": nb,
        "hunks": list(sm.get_grouped_opcodes(args.context)),
        "added": sum(j2 - j1 for t, _, _, j1, j2 in ops if t in ("replace", "insert")),
        "removed": sum(i2 - i1 for t, i1, i2, _, _ in ops if t in ("replace", "delete"))})


def count_changed(res):
    return len(res["changed"]) + len(res["format_changed"]) + len(res["binary_changed"])


def anchors(res):
    """변경 파일 경로 → HTML 리포트 안의 앵커 id (엑셀에서 링크할 때 사용)"""
    return {c["path"]: f"f{i}" for i, c in enumerate(res["changed"])}


def hunk_span(lines, i1, i2):
    """unified diff 헤더용 '시작,줄수' (원래 파일의 줄 번호 기준)"""
    if i1 == i2:
        return f"{lines[i1 - 1][0] if i1 else 0},0"
    start, end = lines[i1][0], lines[i2 - 1][0]
    return str(start) if start == end else f"{start},{end - start + 1}"


def unified_lines(c):
    a, b = c["a"], c["b"]
    yield f"--- A/{c['path']}"
    yield f"+++ B/{c['path']}"
    for group in c["hunks"]:
        yield (f"@@ -{hunk_span(a, group[0][1], group[-1][2])} "
               f"+{hunk_span(b, group[0][3], group[-1][4])} @@")
        for tag, i1, i2, j1, j2 in group:
            if tag == "equal":
                yield from (" " + x[1] for x in a[i1:i2])
                continue
            yield from ("-" + x[1] for x in a[i1:i2])
            yield from ("+" + y[1] for y in b[j1:j2])


def diff_rows(c):
    """나란히 보기용 행 (종류, A 줄번호, A 원문, B 줄번호, B 원문). 종류 sep은 hunk 구분선"""
    a, b = c["a"], c["b"]
    for k, group in enumerate(c["hunks"]):
        if k:
            yield ("sep", None, "", None, "")
        for tag, i1, i2, j1, j2 in group:
            for n in range(max(i2 - i1, j2 - j1)):
                x = a[i1 + n] if i1 + n < i2 else None
                y = b[j1 + n] if j1 + n < j2 else None
                kind = tag if tag == "equal" else "replace" if x and y else "delete" if x else "insert"
                yield (kind, x and x[0], x[1] if x else "", y and y[0], y[1] if y else "")


# ---------------------------------------------------------------- 콘솔 출력
class C:
    RED = GREEN = YELLOW = CYAN = BOLD = RESET = ""


def enable_color():
    if os.name == "nt":
        os.system("")  # Windows 10+ 콘솔에서 ANSI 색상 활성화
    C.RED, C.GREEN, C.YELLOW, C.CYAN = "\033[31m", "\033[32m", "\033[33m", "\033[36m"
    C.BOLD, C.RESET = "\033[1m", "\033[0m"


def print_report(res, dir_a, dir_b, summary_only):
    line = "=" * 70
    print(f"{C.BOLD}{line}\n 폴더 비교\n  A: {dir_a}\n  B: {dir_b}\n{line}{C.RESET}")
    extra = f"  |  무시됨: {len(res['ignored'])}" if res["ignored"] else ""
    extra += f"  |  {C.RED}양쪽에 없음: {len(res['missing'])}{C.RESET}" if res["missing"] else ""
    extra += f"  |  {C.RED}오류: {len(res['errors'])}{C.RESET}" if res["errors"] else ""
    print(f"  동일: {len(res['same'])}  |  변경: {count_changed(res)}"
          f"  |  A에만: {len(res['only_a'])}  |  B에만: {len(res['only_b'])}{extra}\n")

    if res["errors"]:
        print(f"{C.BOLD}[읽기 오류]{C.RESET}")
        for er in res["errors"]:
            print(f"  {C.RED}! {er['path']}{C.RESET}  ({er['detail']})")
        print()
    if res["missing"]:
        print(f"{C.BOLD}[양쪽에 없음 (비교 대상 목록에만 있음)]{C.RESET}")
        for p in res["missing"]:
            print(f"  {C.RED}? {p}{C.RESET}")
        print()
    if res["only_a"]:
        print(f"{C.BOLD}[A에만 있는 파일]{C.RESET}")
        for p in res["only_a"]:
            print(f"  {C.RED}- {p}{C.RESET}")
        print()
    if res["only_b"]:
        print(f"{C.BOLD}[B에만 있는 파일]{C.RESET}")
        for p in res["only_b"]:
            print(f"  {C.GREEN}+ {p}{C.RESET}")
        print()
    if count_changed(res):
        print(f"{C.BOLD}[변경된 파일]{C.RESET}")
        for c in res["changed"]:
            print(f"  {C.YELLOW}* {c['path']}{C.RESET}  ({C.GREEN}+{c['added']}{C.RESET} / {C.RED}-{c['removed']}{C.RESET})")
        for c in res["format_changed"]:
            print(f"  {C.YELLOW}* {c['path']}{C.RESET}  (형식만 다름: {c['detail']})")
        for p in res["binary_changed"]:
            print(f"  {C.YELLOW}* {p}{C.RESET}  (바이너리 파일 - 내용 다름)")
        print()
    if res["ignored"]:
        print(f"{C.BOLD}[공백 차이만 있는 파일 (옵션으로 무시됨)]{C.RESET}")
        for p in res["ignored"]:
            print(f"  {C.CYAN}~ {p}{C.RESET}")
        print()

    if summary_only:
        return
    for c in res["changed"]:
        print(f"{C.BOLD}{'-' * 70}\n {c['path']}\n{'-' * 70}{C.RESET}")
        for l in unified_lines(c):
            if l.startswith("+++") or l.startswith("---"):
                print(f"{C.BOLD}{l}{C.RESET}")
            elif l.startswith("@@"):
                print(f"{C.CYAN}{l}{C.RESET}")
            elif l.startswith("+"):
                print(f"{C.GREEN}{l}{C.RESET}")
            elif l.startswith("-"):
                print(f"{C.RED}{l}{C.RESET}")
            else:
                print(l)
        print()


# ---------------------------------------------------------------- HTML 리포트
HTML_CSS = """
body{font-family:'Segoe UI','Malgun Gothic',sans-serif;margin:24px;color:#222}
h1{font-size:20px} h2{font-size:16px;margin-top:28px;border-bottom:1px solid #ccc;padding-bottom:4px}
.sum span{display:inline-block;margin-right:16px;padding:4px 10px;border-radius:4px;background:#f0f0f0}
ul{font-family:Consolas,monospace;font-size:13px} .add{color:#1a7f37} .del{color:#cf222e}
details{margin:10px 0;border:1px solid #ddd;border-radius:4px} summary{cursor:pointer;padding:6px 10px;background:#f6f8fa;font-family:Consolas,monospace}
table.diff{font-family:Consolas,monospace;font-size:12px;border-collapse:collapse;width:100%;table-layout:fixed}
table.diff td{padding:0 4px;vertical-align:top;white-space:pre-wrap;word-break:break-all}
table.diff col.no{width:4.5em} td.no{background:#eee;color:#888;text-align:right}
tr.delete td.ta{background:#ffebe9} tr.insert td.tb{background:#dafbe1}
tr.replace td.ta,tr.replace td.tb{background:#fff8c5}
tr.replace td.ta mark{background:#ffc1bd} tr.replace td.tb mark{background:#aceebb}
tr.sep td{background:#f6f8fa;color:#888;text-align:center}
"""


def inline_diff(a, b):
    """바뀐 줄 한 쌍에서 달라진 글자를 <mark>로 감싼 HTML 반환"""
    e = html.escape
    if len(a) + len(b) > INLINE_DIFF_MAX:
        return e(a), e(b)
    ha, hb = [], []
    for tag, i1, i2, j1, j2 in difflib.SequenceMatcher(None, a, b, autojunk=False).get_opcodes():
        sa, sb = e(a[i1:i2]), e(b[j1:j2])
        if tag == "equal":
            ha.append(sa)
            hb.append(sb)
            continue
        if sa:
            ha.append(f"<mark>{sa}</mark>")
        if sb:
            hb.append(f"<mark>{sb}</mark>")
    return "".join(ha), "".join(hb)


def html_table(c):
    e = html.escape
    rows = ["<table class='diff'><col class='no'><col><col class='no'><col>"]
    for kind, na, ta, nb, tb in diff_rows(c):
        if kind == "sep":
            rows.append("<tr class='sep'><td colspan='4'>⋯</td></tr>")
            continue
        ha, hb = inline_diff(ta, tb) if kind == "replace" else (e(ta), e(tb))
        rows.append(f"<tr class='{kind}'><td class='no'>{na or ''}</td><td class='ta'>{ha}</td>"
                    f"<td class='no'>{nb or ''}</td><td class='tb'>{hb}</td></tr>")
    rows.append("</table>")
    return "".join(rows)


def write_html(res, dir_a, dir_b, out_path):
    e = html.escape
    parts = [f"<!doctype html><html><head><meta charset='utf-8'><title>폴더 비교 리포트</title>"
             f"<style>{HTML_CSS}</style></head><body>",
             f"<h1>폴더 비교 리포트</h1><p>A: <code>{e(str(dir_a))}</code><br>B: <code>{e(str(dir_b))}</code><br>"
             f"생성: {datetime.now():%Y-%m-%d %H:%M}</p>",
             f"<div class='sum'><span>동일 {len(res['same'])}</span>"
             f"<span>변경 {count_changed(res)}</span>"
             f"<span class='del'>A에만 {len(res['only_a'])}</span>"
             f"<span class='add'>B에만 {len(res['only_b'])}</span>"
             f"<span>무시됨 {len(res['ignored'])}</span>"
             + (f"<span class='del'>양쪽에 없음 {len(res['missing'])}</span>" if res["missing"] else "") +
             f"<span class='del'>오류 {len(res['errors'])}</span></div>"]

    if res["errors"]:
        parts.append("<h2>읽기 오류</h2><ul>" +
                     "".join(f"<li class='del'>{e(er['path'])} — {e(er['detail'])}</li>" for er in res["errors"]) + "</ul>")

    parts.append("<h2>변경 내용 (클릭하여 펼치기)</h2>")
    if not res["changed"]:
        parts.append("<p>내용이 바뀐 텍스트 파일이 없습니다.</p>")
    ids = anchors(res)
    for c in res["changed"]:
        parts.append(f"<details id='{ids[c['path']]}'><summary>{e(c['path'])} "
                     f"<span class='add'>+{c['added']}</span> <span class='del'>-{c['removed']}</span>"
                     f"</summary>{html_table(c)}</details>")

    # 나머지 목록은 엑셀 리포트가 주 용도이므로 접어서 표시
    others = [("양쪽에 없음 (비교 대상 목록에만 있음)", res["missing"], "del"),
              ("A에만 있는 파일", res["only_a"], "del"), ("B에만 있는 파일", res["only_b"], "add"),
              ("형식만 다른 파일 (줄바꿈/인코딩/BOM)",
               [f"{c['path']} — {c['detail']}" for c in res["format_changed"]], ""),
              ("변경된 바이너리 파일", res["binary_changed"], ""),
              ("공백 차이만 있는 파일 (옵션으로 무시됨)", res["ignored"], "")]
    others = [(t, items, cls) for t, items, cls in others if items]
    if others:
        parts.append("<h2>기타 파일 목록</h2>")
        for title, items, cls in others:
            parts.append(f"<details><summary>{e(title)} ({len(items)})</summary><ul>" +
                         "".join(f"<li class='{cls}'>{e(p)}</li>" for p in items) + "</ul></details>")

    # 엑셀 등에서 report.html#f3 처럼 열면 해당 파일을 펼쳐서 보여줌
    parts.append("<script>function openHash(){var d=document.getElementById(location.hash.slice(1));"
                 "if(d&&d.tagName==='DETAILS'){d.open=true;d.scrollIntoView();}}"
                 "addEventListener('hashchange',openHash);openHash();</script></body></html>")
    Path(out_path).write_text("\n".join(parts), encoding="utf-8")


# ---------------------------------------------------------------- 엑셀 리포트
def write_excel(res, dir_a, dir_b, out_path, options, html_path=None):
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.hyperlink import Hyperlink

    fill = {k: PatternFill("solid", fgColor=v) for k, v in {
        "header": "DDDDDD", "red": "FFEBE9", "green": "DAFBE1", "yellow": "FFF8C5",
        "gray": "F0F0F0", "error": "FFC1BD"}.items()}
    bold, mono = Font(bold=True), Font(name="Consolas")

    def put(ws, row, col, value, fill_key=None, font=None):
        if isinstance(value, str):
            value = ILLEGAL_CHARACTERS_RE.sub("", value)[:32767]  # 엑셀 셀 최대 길이
        cell = ws.cell(row, col, value)
        if isinstance(value, str):
            cell.data_type = "s"  # '='로 시작하는 줄이 수식으로 해석되지 않도록
        if fill_key:
            cell.fill = fill[fill_key]
        if font:
            cell.font = font
        return cell

    def header(ws, titles, widths):
        for col, (t, w) in enumerate(zip(titles, widths), 1):
            put(ws, 1, col, t, "header", bold)
            ws.column_dimensions[chr(64 + col)].width = w
        ws.freeze_panes = "A2"

    wb = Workbook()

    # 요약
    ws = wb.active
    ws.title = "요약"
    info = [("A 폴더", str(dir_a)), ("B 폴더", str(dir_b)),
            ("생성", f"{datetime.now():%Y-%m-%d %H:%M}"), ("옵션", options), (None, None),
            ("구분", "개수"), ("동일", len(res["same"])), ("변경", count_changed(res)),
            ("A에만", len(res["only_a"])), ("B에만", len(res["only_b"])),
            ("무시됨", len(res["ignored"])), ("양쪽에 없음", len(res["missing"])),
            ("오류", len(res["errors"]))]
    for r, (k, v) in enumerate(info, 1):
        put(ws, r, 1, k, font=bold)
        put(ws, r, 2, v, "header" if k == "구분" else None, bold if k == "구분" else None)
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 80

    # 파일 목록
    ws = wb.create_sheet("파일 목록")
    header(ws, ["경로", "상태", "추가", "삭제", "비고"], [60, 14, 8, 8, 50])
    rows = ([(c["path"], "변경", c["added"], c["removed"], "", "yellow") for c in res["changed"]] +
            [(c["path"], "형식만 다름", None, None, c["detail"], "yellow") for c in res["format_changed"]] +
            [(p, "바이너리 변경", None, None, "", "yellow") for p in res["binary_changed"]] +
            [(p, "A에만", None, None, "", "red") for p in res["only_a"]] +
            [(p, "B에만", None, None, "", "green") for p in res["only_b"]] +
            [(p, "무시됨", None, None, "옵션으로 무시된 공백 차이", "gray") for p in res["ignored"]] +
            [(p, "양쪽에 없음", None, None, "비교 대상 목록에만 있음", "error") for p in res["missing"]] +
            [(er["path"], "오류", None, None, er["detail"], "error") for er in res["errors"]] +
            [(p, "동일", None, None, "", None) for p in res["same"]])
    # 변경 파일 경로를 HTML 리포트의 해당 diff로 링크 (엑셀 파일 기준 상대 경로)
    link, ids = None, anchors(res)
    if html_path:
        try:
            link = Path(os.path.relpath(Path(html_path).resolve(), Path(out_path).resolve().parent)).as_posix()
        except ValueError:  # 드라이브가 다르면 절대 경로
            link = Path(html_path).resolve().as_uri()
    link_font = Font(name="Consolas", color="0563C1", underline="single")
    for r, (*values, color) in enumerate(rows, 2):
        for col, v in enumerate(values, 1):
            put(ws, r, col, v, color if col == 2 else None, mono if col == 1 else None)
        if link and values[0] in ids and values[1] == "변경":
            cell = ws.cell(r, 1)
            cell.hyperlink = Hyperlink(ref=cell.coordinate, target=link, location=ids[values[0]],
                                       tooltip="HTML 리포트에서 변경 내용 보기")
            cell.font = link_font
    ws.auto_filter.ref = f"A1:E{len(rows) + 1}"

    wb.save(out_path)


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description="두 폴더의 코드를 같은 파일 이름(상대 경로) 기준으로 비교합니다.")
    ap.add_argument("dir_a", help="기준 폴더 (A)")
    ap.add_argument("dir_b", help="비교 폴더 (B)")
    ap.add_argument("-w", "--ignore-whitespace", action="store_true", help="줄 앞뒤 공백 차이 무시")
    ap.add_argument("-B", "--ignore-blank-lines", action="store_true", help="빈 줄 차이 무시")
    ap.add_argument("-x", "--exclude", nargs="+", metavar="PATTERN",
                    help="제외할 패턴 ('/'가 들어가면 상대 경로 전체와 비교)")
    ap.add_argument("-e", "--ext", nargs="+", metavar="EXT", help="비교할 확장자만 지정")
    ap.add_argument("-c", "--context", type=int, default=3, help="diff 앞뒤 줄 수 (기본 3)")
    ap.add_argument("-o", "--report", metavar="NAME",
                    help="NAME.xlsx(파일 목록)와 NAME.html(변경 내용)을 함께 저장")
    ap.add_argument("--html", metavar="FILE", help="HTML 리포트(변경 내용)만 저장")
    ap.add_argument("--excel", metavar="FILE", help="엑셀 리포트(파일 목록)만 저장")
    ap.add_argument("-l", "--list", metavar="FILE",
                    help="비교 대상 목록 파일 (.xlsx는 첫 시트 A열, 그 외는 한 줄에 하나). 목록에 있는 파일만 비교")
    ap.add_argument("--summary", action="store_true", help="요약만 출력")
    ap.add_argument("--no-color", action="store_true", help="콘솔 색상 끄기")
    args = ap.parse_args()

    dir_a, dir_b = Path(args.dir_a).resolve(), Path(args.dir_b).resolve()
    for d in (dir_a, dir_b):
        if not d.is_dir():
            sys.exit(f"폴더를 찾을 수 없습니다: {d}")
    if args.report:
        base = args.report
        if Path(base).suffix.lower() in (".html", ".xlsx"):
            base = base[:-len(Path(base).suffix)]
        args.html = args.html or base + ".html"
        args.excel = args.excel or base + ".xlsx"
    for out in (args.html, args.excel):
        if out:
            Path(out).resolve().parent.mkdir(parents=True, exist_ok=True)
    list_path = Path(args.list) if args.list else None
    if list_path:
        if not list_path.is_file():
            sys.exit(f"비교 대상 목록 파일을 찾을 수 없습니다: {list_path.resolve()}")
        if list_path.suffix.lower() == ".xls":
            sys.exit("구버전 엑셀(.xls)은 읽을 수 없습니다. .xlsx로 저장한 뒤 사용하세요.")
    if args.excel or (list_path and list_path.suffix.lower() in (".xlsx", ".xlsm")):
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            sys.exit("엑셀 파일을 다루려면 openpyxl이 필요합니다: pip install openpyxl")
    targets = load_targets(list_path) if list_path else None
    if targets == []:
        sys.exit(f"비교 대상 목록이 비어 있습니다: {list_path.resolve()}")

    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if not args.no_color and sys.stdout.isatty():
        enable_color()

    res = compare(dir_a, dir_b, args, targets)
    print_report(res, dir_a, dir_b, args.summary)

    if args.html:
        write_html(res, dir_a, dir_b, args.html)
        print(f"HTML 리포트 저장: {Path(args.html).resolve()}")
    if args.excel:
        options = " ".join(a for a in sys.argv[1:] if a not in (args.dir_a, args.dir_b)) or "(없음)"
        write_excel(res, dir_a, dir_b, args.excel, options, args.html)
        print(f"엑셀 리포트 저장: {Path(args.excel).resolve()}")

    # 차이가 있으면 종료 코드 1, 읽기 오류가 있으면 2 (CI·배치 스크립트에서 활용 가능)
    if res["errors"]:
        sys.exit(2)
    has_diff = any(res[k] for k in ("only_a", "only_b", "changed", "format_changed", "binary_changed", "missing"))
    sys.exit(1 if has_diff else 0)


if __name__ == "__main__":
    main()
