#!/usr/bin/env python3
"""
C 함수 단위 비교 도구 (func_diff.py)

이전 버전과 현재 버전 폴더에서 같은 상대 경로의 .c 파일끼리 짝지은 뒤, 파일 안의 함수를 이름으로 짝지어
어떤 함수가 추가·삭제·변경되었는지 보여줍니다.
함수 순서가 바뀌어도 같은 함수끼리 비교하고, 함수 밖 코드(전역 변수·#define 등)는 '(함수 외)'로 묶어 비교합니다.

사용 예:
    python tools/func_diff.py 이전폴더 현재폴더 -l 목록.xlsx -o reports/report   # 목록(파일명·함수명)에 있는 함수만
    python tools/func_diff.py 이전폴더 현재폴더 -o reports/report                # 모든 함수 + (함수 외)

비교 대상 목록 (-l):
    .xlsx는 첫 시트 A열 = 파일명, B열 = 함수명 (파일명이 비어 있거나 병합 셀이면 위 파일명을 이어 씀)
    텍스트는 한 줄에 '파일명,함수명' (탭도 가능, #은 주석)
    함수명이 비어 있으면 그 파일의 모든 함수('(함수 외)' 제외), '(함수 외)'는 함수 밖 영역(명시한 경우에만 포함)
    같은 이름 함수가 여러 개면 'f'는 모두, 'f#2'는 두 번째 함수만

그 밖의 옵션(-w -B -x -e -c -o --html --excel --summary --no-color)은 folder_diff.py와 같습니다.
종료 코드: 0 = 차이 없음, 1 = 차이 있음, 2 = 읽기 오류 또는 분석 불가 파일 있음
folder_diff.py와 같은 폴더에 있어야 합니다.
"""
import argparse
import bisect
import html
import re
import sys
from datetime import datetime
from pathlib import Path

import folder_diff as fd

FUNC_OUTSIDE = "(함수 외)"   # 함수에 속하지 않는 줄을 묶은 단위
WHOLE_FILE = "(파일 전체)"   # 분석 불가 파일이나 목록의 '모든 함수' 항목
NOT_FUNC_NAMES = {"__attribute__", "__declspec", "if", "while", "for", "switch", "sizeof", "return"}
IDENT_END = re.compile(r"([A-Za-z_]\w*)\s*$")
ALL_CAPS = re.compile(r"[A-Z_][A-Z0-9_]*")


# ---------------------------------------------------------------- C 스캐너
def mask_c(text):
    """주석·문자열·문자 리터럴·전처리 줄을 공백으로 가린 같은 길이의 문자열 (줄바꿈은 유지)"""
    out = list(text)
    n, i, line_start = len(text), 0, True

    def blank(a, b):
        for k in range(a, b):
            if out[k] != "\n":
                out[k] = " "

    while i < n:
        ch = text[i]
        if ch == "\n":
            line_start = True
            i += 1
            continue
        if line_start and ch in " \t":
            i += 1
            continue
        if line_start and ch == "#":  # 전처리 줄: 주석·문자열을 건너뛰며 논리적 줄 끝('\' 이어짐 포함)까지
            j = i
            while j < n:
                c = text[j]
                if c == "\n":
                    k = j - 1
                    while k >= i and text[k] in " \t\r":
                        k -= 1
                    if k >= i and text[k] == "\\":
                        j += 1
                        continue
                    break
                if text.startswith("/*", j):
                    e = text.find("*/", j + 2)
                    j = n if e < 0 else e + 2
                elif text.startswith("//", j):
                    e = text.find("\n", j)
                    j = n if e < 0 else e
                elif c in "\"'":
                    j += 1
                    while j < n and text[j] != c and text[j] != "\n":
                        j += 2 if text[j] == "\\" else 1
                    if j < n and text[j] == c:
                        j += 1
                else:
                    j += 1
            j = min(j, n)
            blank(i, j)
            i = j
            continue
        line_start = False
        if text.startswith("/*", i):
            e = text.find("*/", i + 2)
            e = n if e < 0 else e + 2
            blank(i, e)
            i = e
        elif text.startswith("//", i):
            e = text.find("\n", i)
            e = n if e < 0 else e
            blank(i, e)
            i = e
        elif ch in "\"'":
            j = i + 1
            while j < n and text[j] != ch and text[j] != "\n":
                j += 2 if text[j] == "\\" else 1
            e = min(j + 1, n) if j < n and text[j] == ch else min(j, n)
            blank(i, e)
            i = e
        else:
            i += 1
    return "".join(out)


def func_name(header):
    """최상위 '{' 앞의 선언부(가려진 텍스트)가 함수 정의면 함수 이름, 아니면 None"""
    h = header.rstrip()

    # 뒤에 붙은 __attribute__(…) 또는 __declspec(…) 제거
    while True:
        if not h.endswith(")"):
            return None
        # 마지막 ) 앞의 ( 찾기
        depth = 0
        for k in range(len(h) - 1, -1, -1):
            if h[k] == ")":
                depth += 1
            elif h[k] == "(":
                depth -= 1
                if depth == 0:
                    break
        else:
            return None
        # k부터 len(h)까지가 마지막 (...) 그룹
        paren_group = h[k:len(h)]
        rest = h[:k].rstrip()
        # rest의 끝이 __attribute__ 또는 __declspec인지 확인
        if rest.endswith("__attribute__") or rest.endswith("__declspec"):
            # 그 부분을 제거
            if rest.endswith("__attribute__"):
                h = rest[:-len("__attribute__")].rstrip()
            else:
                h = rest[:-len("__declspec")].rstrip()
            continue
        # 더 이상 제거할 것이 없으면 나감
        break

    if not h.endswith(")"):
        return None
    depth = 0
    for k in range(len(h) - 1, -1, -1):  # 마지막 ( … ) 묶음의 여는 괄호 찾기
        if h[k] == ")":
            depth += 1
        elif h[k] == "(":
            depth -= 1
            if depth == 0:
                break
    else:
        return None
    depth = 0
    for ch in h:  # 괄호 밖 '='가 있으면 초기화 식 (int a[] = {…})
        depth += (ch == "(") - (ch == ")")
        if ch == "=" and depth == 0:
            return None
    m = IDENT_END.search(h[:k])
    if not m or m.group(1) in NOT_FUNC_NAMES:
        return None
    if ALL_CAPS.fullmatch(h[:k].strip()):  # TASK(T10ms), ISR(Can_Isr): 리턴 타입 없는 OS 매크로면 괄호 안이 이름
        inner = re.fullmatch(r"\(\s*([A-Za-z_]\w*)\s*\)", h[k:])
        if inner:
            return inner.group(1)
    return m.group(1)


def find_functions(lines):
    """줄 목록에서 최상위 함수 정의를 찾아 [(이름, 시작 줄, 끝 줄)] 반환 (줄 번호는 1부터).
    중괄호 짝이 맞지 않으면 None"""
    m = mask_c("\n".join(lines))
    starts = [0] + [x.end() for x in re.finditer("\n", m)]
    funcs, depth, seg, head = [], 0, 0, None
    for x in re.finditer(r"[{};]", m):
        i, ch = x.start(), x.group()
        if ch == "{":
            if depth == 0:
                head = (seg, i)
            depth += 1
        elif ch == "}":
            depth -= 1
            if depth < 0:
                return None
            if depth == 0:
                text = m[head[0]:head[1]]
                name = func_name(text)
                if name:
                    first = head[0] + len(text) - len(text.lstrip())
                    funcs.append((name, bisect.bisect_right(starts, first), bisect.bisect_right(starts, i)))
                seg = i + 1
        elif depth == 0:  # ';'
            seg = i + 1
    return funcs if depth == 0 else None


def split_units(lines):
    """파일을 비교 단위로 나눔: {이름: [줄 번호…]}. 첫 항목은 '(함수 외)'(함수에 속하지 않는 줄),
    그 뒤로 함수가 파일 순서대로 옴. 같은 이름이 또 나오면 '이름#2'. 분석할 수 없으면 None"""
    funcs = find_functions(lines)
    if funcs is None:
        return None
    units, seen, inside = {FUNC_OUTSIDE: []}, {}, set()
    for name, s, e in funcs:
        seen[name] = seen.get(name, 0) + 1
        units[name if seen[name] == 1 else f"{name}#{seen[name]}"] = list(range(s, e + 1))
        inside.update(range(s, e + 1))
    units[FUNC_OUTSIDE] = [n for n in range(1, len(lines) + 1) if n not in inside]
    return units


# ---------------------------------------------------------------- 비교 대상 목록
FUNC_HEADERS = {"함수", "함수명", "함수 이름", "함수이름", "function", "func", "function name"}


def load_func_list(path):
    """목록 파일 → {파일 항목: 함수명 집합 또는 None(그 파일의 모든 함수)} (목록 순서 유지)"""
    plan, cur = {}, None
    first = True  # 아직 비어 있지 않은 행을 못 봄
    for f, fn in fd.read_rows(path, 2):
        f = "" if f is None else str(f).strip()
        fn = "" if fn is None else str(fn).strip()
        if not (f or fn):
            continue
        was_first, first = first, False
        if was_first and (f.casefold() in fd.TARGET_HEADERS or fn.casefold() in FUNC_HEADERS):
            continue  # 제목 행 (처음 나오는 비어 있지 않은 행)
        cur = f or cur  # 파일명이 비어 있거나 병합 셀이면 위 파일명을 이어 씀
        if not cur:
            continue
        if fn.replace(" ", "") in ("(함수외)", "함수외"):
            fn = FUNC_OUTSIDE
        if not fn:
            plan[cur] = None
        elif cur not in plan:
            plan[cur] = {fn}
        elif plan[cur] is not None:
            plan[cur].add(fn)
    return plan


# ---------------------------------------------------------------- 함수 단위 비교
SAME, CHANGED, DELETED, ADDED = "동일", "변경", "삭제", "추가"
NO_FUNC, NO_FILE, BROKEN = "함수를 찾을 수 없음", "파일을 찾을 수 없음", "분석 불가"


def make_row(file, func, status, a_range="", b_range="", diff=None, note=""):
    return {"file": file, "func": func, "status": status, "a_range": a_range, "b_range": b_range,
            "added": diff["added"] if diff else None, "removed": diff["removed"] if diff else None,
            "note": note, "diff": diff}


def read_side(path, args):
    """파일 한쪽을 읽어 (split_units 결과, normalize 결과, {줄 번호: normalize 항목}) 반환. 파일이 없으면 None"""
    if path is None:
        return None
    lines, _ = fd.decode_lines(path.read_bytes())
    norm = fd.normalize(lines, args.ignore_whitespace, args.ignore_blank_lines)
    return split_units(lines), norm, {x[0]: x for x in norm}


def pick(side, nums):
    """nums 줄만 골라냄 (-B로 빠진 빈 줄은 색인에 없으므로 제외됨)"""
    if side is None or nums is None:
        return []
    idx = side[2]
    return [idx[n] for n in nums if n in idx]


def line_range(name, nums):
    return f"{nums[0]}-{nums[-1]}" if nums and name != FUNC_OUTSIDE else ""


def compare_c_file(rel, pa, pb, wanted, args, rows, all_funcs=False):
    """C 파일 한 쌍을 함수 단위로 비교해 rows에 결과 행을 추가. wanted: None(전부, (함수 외) 포함) 또는
    함수명 집합('f#2'처럼 번호 붙은 이름도 가능). all_funcs: 목록의 '모든 함수' 항목 -
    (함수 외)는 wanted에 명시한 경우에만 포함"""
    sa, sb = read_side(pa, args), read_side(pb, args)
    broken = [s for s, side in (("이전", sa), ("현재", sb)) if side and side[0] is None]
    if broken:  # 중괄호 짝이 안 맞으면 파일 전체 diff로 대체
        d = fd.line_diff(sa[1] if sa else [], sb[1] if sb else [], args.context)
        rows.append(make_row(rel, WHOLE_FILE, BROKEN, diff=d,
                             note=f"중괄호 짝 불일치({', '.join(broken)}) - 파일 전체 비교로 대체"))
        return
    ua, ub = (sa[0] if sa else {}), (sb[0] if sb else {})
    names = list(ua) + [k for k in ub if k not in ua]
    if wanted is not None:
        names = [k for k in names
                 if (all_funcs and k != FUNC_OUTSIDE) or k in wanted or k.split("#")[0] in wanted]
    for name in names:
        la, lb = ua.get(name), ub.get(name)
        na, nb = pick(sa, la), pick(sb, lb)
        if name == FUNC_OUTSIDE and not na and not nb:
            continue
        d = fd.line_diff(na, nb, args.context)
        if la is not None and lb is not None:
            status = CHANGED if d else SAME
        else:  # 이전에만 있으면 삭제, 현재에만 있으면 추가 (본문 전체가 diff로 보임)
            status = DELETED if la is not None else ADDED
        rows.append(make_row(rel, name, status, line_range(name, la), line_range(name, lb), d))
    if wanted is not None:
        found = set(names) | {k.split("#")[0] for k in names}
        for fn in sorted(wanted - found):
            rows.append(make_row(rel, fn, NO_FUNC, note="이전·현재 버전 파일 모두에 없음 (함수명은 대소문자 구분)"))


def compare_funcs(dir_a, dir_b, args, plan=None):
    """두 폴더의 C 파일을 함수 단위로 비교. plan은 load_func_list() 결과(없으면 전체 비교).
    반환: {"rows": [결과 행…], "errors": [{"path", "detail"}…]}"""
    excludes, exts = fd.file_filters(args)
    errors, rows = [], []
    fa = fd.collect_files(dir_a, excludes, exts, errors)
    fb = fd.collect_files(dir_b, excludes, exts, errors)
    if plan is None:
        targets = {rel: (None, False) for rel in set(fa) | set(fb)}
    else:
        targets, every = {}, set(fa) | set(fb)
        for entry, funcs in plan.items():
            hits, _ = fd.match_targets(every, [entry])
            if not hits:
                for fn in (sorted(funcs) if funcs else [WHOLE_FILE]):
                    rows.append(make_row(entry, fn, NO_FILE, note="이전·현재 버전 폴더 모두에 없음"))
                continue
            for rel in hits:  # 같은 파일이 목록에 여러 번 나오면 함수 집합을 합침
                prev, every_func = targets.get(rel, (set(), False))
                targets[rel] = (prev | (funcs or set()), every_func or funcs is None)
    for rel in sorted(targets):
        wanted, all_funcs = targets[rel]
        if all_funcs:  # 모든 함수가 대상이면 개별 함수명은 의미 없고 (함수 외)만 따로 판단
            wanted = wanted & {FUNC_OUTSIDE}
        try:
            compare_c_file(rel, fa.get(rel), fb.get(rel), wanted, args, rows, all_funcs)
        except OSError as ex:
            errors.append({"path": rel, "detail": f"파일 읽기 실패: {ex.strerror or ex}"})
    return {"rows": rows, "errors": errors}


# ---------------------------------------------------------------- 출력
STATUS_ORDER = [ADDED, DELETED, CHANGED, SAME, NO_FUNC, NO_FILE, BROKEN]
STATUS_FILL = {CHANGED: "yellow", DELETED: "red", ADDED: "green",
               NO_FUNC: "error", NO_FILE: "error", BROKEN: "error", SAME: None}


def counts(res):
    c = {s: 0 for s in STATUS_ORDER}
    for r in res["rows"]:
        c[r["status"]] += 1
    return c


def print_report(res, dir_a, dir_b, summary_only):
    C = fd.C
    line = "=" * 70
    print(f"{C.BOLD}{line}\n 함수 변경 비교\n  이전: {dir_a}\n  현재: {dir_b}\n{line}{C.RESET}")
    summary = [f"{s}: {n}" for s, n in counts(res).items() if n or s in (ADDED, DELETED, CHANGED, SAME)]
    errors = f"  |  {C.RED}오류: {len(res['errors'])}{C.RESET}" if res["errors"] else ""
    print("  " + "  |  ".join(summary) + errors + "\n")
    for er in res["errors"]:
        print(f"  {C.RED}! {er['path']}{C.RESET}  ({er['detail']})")
    color = {CHANGED: C.YELLOW, DELETED: C.RED, ADDED: C.GREEN,
             NO_FUNC: C.RED, NO_FILE: C.RED, BROKEN: C.RED, SAME: ""}
    current = None
    for r in res["rows"]:
        if r["file"] != current:
            current = r["file"]
            print(f"{C.BOLD}[{current}]{C.RESET}")
        stat = f" (+{r['added']} / -{r['removed']})" if r["status"] == CHANGED else ""
        note = f"  - {r['note']}" if r["note"] else ""
        print(f"  {color[r['status']]}{r['status']}{C.RESET}  {r['func']}{stat}{note}")
    print()
    if summary_only:
        return
    for r in res["rows"]:
        if r["diff"]:
            print(f"{C.BOLD}{'-' * 70}\n {r['file']} :: {r['func']}  [{r['status']}]\n{'-' * 70}{C.RESET}")
            fd.print_unified({"path": f"{r['file']}::{r['func']}", **r["diff"]})
            print()


def write_html(res, dir_a, dir_b, out_path):
    e = html.escape
    parts = [f"<!doctype html><html><head><meta charset='utf-8'><title>함수 변경 리포트</title>"
             f"<style>{fd.HTML_CSS}</style></head><body>",
             f"<h1>함수 변경 리포트</h1><p>이전: <code>{e(str(dir_a))}</code><br>현재: <code>{e(str(dir_b))}</code><br>"
             f"생성: {datetime.now():%Y-%m-%d %H:%M}</p>",
             "<div class='sum'>" + "".join(
                 f"<span class='{ {ADDED: 'add', DELETED: 'del'}.get(s, '') }'>{e(s)} {n}</span>"
                 for s, n in counts(res).items())
             + f"<span class='del'>오류 {len(res['errors'])}</span></div>"]
    if res["errors"]:
        parts.append("<h2>읽기 오류</h2><ul>" + "".join(
            f"<li class='del'>{e(er['path'])} — {e(er['detail'])}</li>" for er in res["errors"]) + "</ul>")
    missing = [r for r in res["rows"] if r["status"] in (NO_FUNC, NO_FILE)]
    if missing:
        parts.append("<h2>목록에서 찾지 못한 항목</h2><ul>" + "".join(
            f"<li class='del'>{e(r['file'])} :: {e(r['func'])} — {e(r['status'])}</li>" for r in missing) + "</ul>")
    by_file = {}
    for k, r in enumerate(res["rows"]):
        if r["status"] not in (NO_FUNC, NO_FILE):
            by_file.setdefault(r["file"], []).append((k, r))
    for file, items in by_file.items():
        parts.append(f"<h2>{e(file)}</h2>")
        for k, r in items:
            if r["diff"]:
                stat = f" <span class='add'>+{r['added']}</span> <span class='del'>-{r['removed']}</span>"
                parts.append(f"<details id='r{k}'><summary>{e(r['func'])} [{e(r['status'])}]{stat}"
                             f"</summary>{fd.html_table(r['diff'])}</details>")
        rest = [r for _, r in items if not r["diff"]]
        if rest:
            parts.append(f"<details><summary>차이 없는 항목 ({len(rest)})</summary><ul>" + "".join(
                f"<li>{e(r['func'])} [{e(r['status'])}]</li>" for r in rest) + "</ul></details>")
    parts.append(fd.HASH_SCRIPT + "</body></html>")
    Path(out_path).write_text("\n".join(parts), encoding="utf-8")


def write_excel(res, dir_a, dir_b, out_path, options, html_path=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    mono = Font(name="Consolas")
    wb = Workbook()
    ws = wb.active
    ws.title = "요약"
    fd.xl_summary_sheet(ws, dir_a, dir_b, options,
                        list(counts(res).items()) + [("오류", len(res["errors"]))])

    ws = wb.create_sheet("함수 목록")
    fd.xl_header(ws, ["파일", "함수", "상태", "이전 줄", "현재 줄", "추가 줄", "삭제 줄", "비고", "검토 결과", "검토 의견"],
                 [40, 32, 18, 12, 12, 8, 8, 40, 12, 50])
    link = fd.relative_link(html_path, out_path) if html_path else None
    rows = [(r["file"], r["func"], r["status"], r["a_range"], r["b_range"], r["added"], r["removed"],
             r["note"], STATUS_FILL[r["status"]], f"r{k}" if r["diff"] else None)
            for k, r in enumerate(res["rows"])]
    rows += [(er["path"], "", "오류", "", "", None, None, er["detail"], "error", None) for er in res["errors"]]
    for n, (*values, fill_key, anchor) in enumerate(rows, 2):
        for col, v in enumerate(values, 1):
            fd.xl_put(ws, n, col, v, fill_key if col == 3 else None, mono if col in (1, 2) else None)
        if link and anchor:
            fd.xl_link(ws.cell(n, 2), link, anchor)
    ws.auto_filter.ref = f"A1:J{len(rows) + 1}"
    wb.save(out_path)


def exit_code(res):
    if res["errors"] or any(r["status"] == BROKEN for r in res["rows"]):
        return 2
    return 1 if any(r["status"] != SAME for r in res["rows"]) else 0


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description="이전 버전과 현재 버전의 C 파일을 함수 단위로 비교해 추가·삭제·변경된 함수를 보여줍니다.")
    fd.add_common_args(ap, ext_default=[".c"])
    ap.add_argument("-l", "--list", metavar="FILE",
                    help="비교 대상 목록 (.xlsx: A열 파일명·B열 함수명 / 텍스트: 한 줄에 '파일명,함수명')")
    args = ap.parse_args()

    dir_a, dir_b = fd.check_dirs(args)
    fd.resolve_report_paths(args)
    list_path = fd.check_list_path(args.list)
    fd.require_openpyxl(args, list_path)
    plan = load_func_list(list_path) if list_path else None
    if plan == {}:
        sys.exit(f"비교 대상 목록이 비어 있습니다: {list_path.resolve()}")
    fd.setup_console(args.no_color)

    res = compare_funcs(dir_a, dir_b, args, plan)
    print_report(res, dir_a, dir_b, args.summary)
    if args.html:
        write_html(res, dir_a, dir_b, args.html)
        print(f"HTML 리포트 저장: {Path(args.html).resolve()}")
    if args.excel:
        write_excel(res, dir_a, dir_b, args.excel, fd.options_text(args), args.html)
        print(f"엑셀 리포트 저장: {Path(args.excel).resolve()}")
    sys.exit(exit_code(res))


if __name__ == "__main__":
    main()
